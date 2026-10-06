from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.companies.shemas import CompanyRole

COMPANY_NAME_HEADER = "Наименование"
COMPANY_INN_HEADER = "ИНН"
DIRECTOR_NAME_HEADER = "ФИО руководителя"
DIRECTOR_INN_HEADER = "ИННФЛ руководителя"
DIRECTOR_POSITION_HEADER = "Должность руководителя"
PRIMARY_PHONE_HEADER = "Номер телефона"
PRIMARY_EMAIL_HEADER = "Электронная почта"
WEBSITE_HEADER = "Ссылка на сайт"
SOURCE_HEADER = "Источник"
SEGMENT_HEADER = "Название сегмента"

REQUIRED_HEADERS = {COMPANY_NAME_HEADER, COMPANY_INN_HEADER}
MAX_IMPORT_ROWS = 10_000


class CompanyExcelFormatError(ValueError):
    pass


@dataclass
class ParsedCompanyRow:
    row_number: int
    data: dict[str, Any]


def _as_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = str(value).strip()
    return result or None


def _unique(values: list[str | None]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        if not value:
            continue
        key = value.casefold()
        if key not in seen:
            seen.add(key)
            result.append(value)
    return result


def _find_data_sheet(workbook):
    for sheet in workbook.worksheets:
        headers = [_as_text(cell.value) for cell in sheet[1]]
        if REQUIRED_HEADERS.issubset(set(headers)):
            return sheet, headers
    raise CompanyExcelFormatError(
        "В Excel не найдены обязательные колонки: Наименование и ИНН",
    )


def parse_companies_excel(
        content: bytes,
        role: CompanyRole | None = None,
) -> tuple[list[ParsedCompanyRow], list[dict]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise CompanyExcelFormatError("Не удалось прочитать Excel-файл") from exc

    try:
        sheet, headers = _find_data_sheet(workbook)
        header_indexes = {
            header: index
            for index, header in enumerate(headers)
            if header is not None
        }
        phone_indexes = [
            index for index, header in enumerate(headers)
            if header == PRIMARY_PHONE_HEADER or (header or "").startswith("Дополнительный телефон")
        ]
        email_indexes = [
            index for index, header in enumerate(headers)
            if header == PRIMARY_EMAIL_HEADER or (header or "").startswith("Дополнительная электронная почта")
        ]

        parsed_rows: list[ParsedCompanyRow] = []
        errors: list[dict] = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if row_number > MAX_IMPORT_ROWS + 1:
                raise CompanyExcelFormatError(
                    f"В файле больше {MAX_IMPORT_ROWS} строк данных",
                )
            if not any(value not in (None, "") for value in values):
                continue

            def get(header: str) -> str | None:
                index = header_indexes.get(header)
                return _as_text(values[index]) if index is not None and index < len(values) else None

            name = get(COMPANY_NAME_HEADER)
            inn = get(COMPANY_INN_HEADER)
            if not name or not inn:
                errors.append({
                    "row": row_number,
                    "detail": "Не заполнены обязательные поля Наименование или ИНН",
                })
                continue

            phones = _unique([
                _as_text(values[index]) if index < len(values) else None
                for index in phone_indexes
            ])
            emails = _unique([
                _as_text(values[index]) if index < len(values) else None
                for index in email_indexes
            ])
            website = get(WEBSITE_HEADER)
            segment = get(SEGMENT_HEADER)

            data: dict[str, Any] = {
                "name": name,
                "inn": inn,
                "phones": phones,
                "emails": emails,
                "websites": [website] if website else [],
                "source": get(SOURCE_HEADER),
                "segments": [segment] if segment else [],
            }
            if role:
                data["roles"] = [role]

            director_name = get(DIRECTOR_NAME_HEADER)
            if director_name:
                contact: dict[str, Any] = {
                    "name": director_name,
                    "isPrimary": True,
                }
                director_inn = get(DIRECTOR_INN_HEADER)
                director_position = get(DIRECTOR_POSITION_HEADER)
                if director_inn:
                    contact["inn"] = director_inn
                if director_position:
                    contact["position"] = director_position
                data["contactPersons"] = [contact]

            parsed_rows.append(ParsedCompanyRow(row_number=row_number, data=data))

        return parsed_rows, errors
    finally:
        workbook.close()
